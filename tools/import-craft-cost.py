r"""生產成本資料：把站長提供的 Google 試算表（生產成本計算）轉成 static/data/crafting-cost.json
用法：python tools/import-craft-cost.py <試算表匯出的 xlsx>
  另外會抓攻略網「最佳採集點、特殊材料掉落點」https://guide.starcg.net/production/best-spots 的素材屋表：
  每種採集材料自己的取出費（4000 個）、等級、素材屋採滿時間（「-」＝素材屋沒有，用同等級的預設時間）
  （匯出：https://docs.google.com/spreadsheets/d/1Lp6ewdWI9EKelNypD6JoyvbBkU18eDfk_VXLiMux9Kc/export?format=xlsx）
試算表的算法（Home 分頁）：
  材料成本 ＝ 取出費用（採集屋一次取 4000 個）÷ 4000 × 數量；礦條 1 條＝20 個礦，取出費用 ×20
  工時成本 ＝ 數量 ÷ 4000 × 採滿時間(分) ÷ 60 × 工時價格(魔幣/小時) × 倍率（一般 1、礦條 20）
  特殊材料、布：直接用每件魔幣；每件總成本 ＝ 材料 ＋ 工時 ＋ 製作費；料理／藥水一箱 60 件另加入箱 800
"""
import io, json, os, re, sys, html, urllib.request
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'static', 'data', 'crafting-cost.json')
# 試算表用的字 → 市場（網站名錄）用的字
VAR = [('鎖鍊', '鎖鏈'), ('蜥蝪', '蜥蜴'), ('蔓陀羅', '曼陀羅'), ('禦', '御'), ('製', '制'), ('癒', '愈'), ('鉤', '鈎'),
       ('鏽', '銹'), ('兇', '凶'), ('著', '着'), ('薑', '姜'), ('蘋', '苹'), ('麵', '面'), ('鬥', '斗'), ('蔔', '卜'), ('密瓜', '蜜瓜')]
# 試算表的簡稱 → 網站配方名稱
def ALIAS(n):
    m = re.match(r'^(\d+)水$', n)
    if m: return f'生命力回復藥({m.group(1)})'
    return {'水龍之藥': '水龍藥水', '香水：深藍九號': '香水深藍九號', '月光之弓+': '月光長弓+'}.get(n, n)
conv = lambda s: __import__('functools').reduce(lambda a, p: a.replace(*p), VAR, str(s).strip())

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
home = wb['Home']
cell = lambda a: home[a].value
levels = [{'take': int(cell(f'C{r}')), 'min': int(cell(f'H{r}'))} for r in range(3, 13)]   # 等級 1～10
labor = float(cell('C18') or 4375)

# 每件價格（魔幣）：Home B21:B36 特殊材料、B50:B61 布；空白＝試算表沒填價格
price = {}
for r in list(range(21, 37)) + list(range(50, 62)):
    n = home[f'B{r}'].value
    if not n or re.match(r'^\d+\(條\)$', str(n)): continue
    v = home[f'C{r}'].value
    price[conv(n)] = float(v) if isinstance(v, (int, float)) else None

# 試算表沒列、但站長確認可以自己採的（礦條的礦名）：鋁條、達馬斯礦條
EXTRA_GATHER = {'鋁', '達馬斯礦'}
# 採集屋採得到的材料（「材料等級」分頁有數字等級的）；其他材料都算「特殊材料」，價格讓使用者自己填
gather = sorted(EXTRA_GATHER | {conv(r[0]) for r in wb['材料等級'].iter_rows(values_only=True) if r and r[0] and isinstance(r[1], (int, float))})
# 布（試算表給的每件價格）
cloth = {n: v for n, v in price.items() if v is not None and n not in ('砂糖', '鋼騎之礦', '魔族的水晶', '妖草的血', '神眼')}
# 特殊材料官方公定價（遊戲內商店，站長 2026-10-27 提供截圖）；沒列到的預設 0
OFFICIAL = {'寶石裝飾膠': 350, '鋼騎之礦': 24000, '魔族的水晶': 60000, '誓言之證': 24000, '火焰之魂': 24000,
            '德特家的布': 40000, '曼陀羅草的皮': 12000, '風龍蜥的甲殼': 12000, '妖草的血': 6000}
LABOR_CLOCK = 4375   # 工時小鬧鐘(1H) 公定價

# 攻略網素材屋表（比試算表的「每級一個時間」準）
def guide_gather():
    req = urllib.request.Request('https://guide.starcg.net/production/best-spots', headers={'User-Agent': 'Mozilla/5.0'})
    t = urllib.request.urlopen(req, timeout=30).read().decode('utf-8')
    t = re.sub(r'<script.*?</script>|<style.*?</style>', '', t, flags=re.S)
    t = html.unescape(re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', '|', t)))
    t = re.sub(r'(\|\s*)+', '|', t)
    out = {}
    for n, fee, lv, mins in re.findall(r'\|([^|]{1,12})\|([\d,]+)G\|(\d+)\|素材屋\|(?:用時 (\d+) 分鐘|-)\|', t):
        out[n.strip()] = {'fee': int(fee.replace(',', '')), 'lv': int(lv), 'min': int(mins) if mins else None}
    return out
mats = guide_gather()
assert len(mats) > 50, '攻略網素材屋表抓不到，版面可能改了'
gather = sorted(set(gather) | set(mats))

# 製作費：「材料」分頁每個配方的第 2 欄
fee, cat = {}, None
for row in wb['材料'].iter_rows(values_only=True):
    if not row or not row[0]: continue
    name = str(row[0]).strip()
    if row[1] == '生產費用' or (row[1] is None and all(v is None for v in row[2:])): cat = name; continue
    if isinstance(row[1], (int, float)): fee[ALIAS(conv(name))] = row[1]
# 料理、藥水一箱 60 件、入箱費 800（Home I25／I26）
data = {
    'src': 'https://docs.google.com/spreadsheets/d/1Lp6ewdWI9EKelNypD6JoyvbBkU18eDfk_VXLiMux9Kc',
    'labor': LABOR_CLOCK, 'levels': levels, 'barMult': 20, 'box': {'n': 60, 'fee': 800},
    'gather': gather, 'mats': mats, 'cloth': cloth, 'special': OFFICIAL, 'fee': fee,
    # 試算表沒列到的配方：依等級套用（試算表裡同等級的製作費都一樣：裝備 40～2400、料理／藥水 20～200）
    # 攻略網「製作價格表」：裝備 1～11 級、補品 1～11 級
    'feeByLv': {'equip': [40, 80, 150, 250, 400, 600, 880, 1260, 1760, 2400, 3400], 'food': [20 * i for i in range(1, 12)], 'potion': [20 * i for i in range(1, 12)]},
}
json.dump(data, io.open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('levels', len(levels), 'mats', len(mats), 'cloth', len(cloth), 'fee', len(fee), '→', OUT)
